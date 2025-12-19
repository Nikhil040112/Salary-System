from fastapi import (
    FastAPI, Request, Depends, Form,
    UploadFile, File, status
)
from fastapi.responses import (
    HTMLResponse, RedirectResponse, FileResponse
)
from fastapi.templating import Jinja2Templates
from sqlalchemy.orm import Session
from openpyxl import load_workbook
from datetime import date, datetime, timedelta, time
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
import os, re

from .database import Base, engine, SessionLocal
from .models import Employee, Penalty, Attendance, SalaryRun
from .auth import verify_login, login_required, serializer

# ---------------- INIT ----------------

Base.metadata.create_all(bind=engine)

app = FastAPI()
templates = Jinja2Templates(directory="app/templates")

# ---------------- DB DEP ----------------

def get_db():
    db = SessionLocal()
    try:
        yield db
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()

# ---------------- LOGIN ----------------

@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request):
    return templates.TemplateResponse("login.html", {"request": request})


@app.post("/login")
def login(
    request: Request,
    username: str = Form(...),
    password: str = Form(...)
):
    if not verify_login(username, password):
        return templates.TemplateResponse(
            "login.html",
            {"request": request, "error": "Invalid credentials"}
        )

    response = RedirectResponse("/", status_code=status.HTTP_303_SEE_OTHER)
    token = serializer.dumps({"user": username})
    response.set_cookie(
        "session",
        token,
        httponly=True,
        samesite="lax"
    )
    return response


@app.get("/logout")
def logout():
    response = RedirectResponse("/login", status_code=303)
    response.delete_cookie("session")
    return response

# ---------------- DASHBOARD ----------------

@app.get("/", response_class=HTMLResponse)
def dashboard(
    request: Request,
    db: Session = Depends(get_db),
    auth=Depends(login_required)
):
    if auth:
        return auth

    today = date.today()
    month_start = date(today.year, today.month, 1)

    return templates.TemplateResponse(
        "dashboard.html",
        {
            "request": request,
            "active_employees": db.query(Employee)
                .filter((Employee.is_active == True) | (Employee.is_active.is_(None)))
                .count(),
            "inactive_employees": db.query(Employee)
                .filter(Employee.is_active == False)
                .count(),
            "penalties_this_month": db.query(Penalty)
                .filter(Penalty.date.between(month_start, today))
                .count()
        }
    )

# ---------------- EMPLOYEES ----------------

@app.get("/employees", response_class=HTMLResponse)
def employee_list(
    request: Request,
    db: Session = Depends(get_db),
    auth=Depends(login_required)
):
    if auth:
        return auth

    employees = db.query(Employee).order_by(Employee.id.asc()).all()
    return templates.TemplateResponse(
        "employee.html",
        {"request": request, "employees": employees}
    )


@app.get("/employees/add", response_class=HTMLResponse)
def add_employee_form(
    request: Request,
    auth=Depends(login_required)
):
    if auth:
        return auth
    return templates.TemplateResponse("add_employee.html", {"request": request})


@app.post("/employees/add")
def add_employee(
    emp_code: str = Form(...),
    name: str = Form(...),
    monthly_salary: float = Form(...),
    salary_cycle_start_day: int = Form(...),
    db: Session = Depends(get_db),
    auth=Depends(login_required)
):
    if auth:
        return auth

    # Prevent duplicate employee code
    if db.query(Employee).filter(Employee.emp_code == emp_code).first():
        return templates.TemplateResponse(
            "add_employee.html",
            {
                "request": request,
                "error": "Employee code already exists"
            }
        )

    employee = Employee(
        emp_code=emp_code,
        name=name,
        monthly_salary=monthly_salary,
        per_day_salary=monthly_salary / 26,
        salary_cycle_start_day=min(salary_cycle_start_day, 28),
        is_active=True
    )
    db.add(employee)
    db.commit()

    return RedirectResponse("/employees", status_code=303)


@app.get("/employees/edit/{employee_id}", response_class=HTMLResponse)
def edit_employee_form(
    request: Request,
    employee_id: int,
    db: Session = Depends(get_db),
    auth=Depends(login_required)
):
    if auth:
        return auth

    employee = db.get(Employee, employee_id)
    if not employee:
        return RedirectResponse("/employees", status_code=303)

    return templates.TemplateResponse(
        "edit_employee.html",
        {"request": request, "employee": employee}
    )


@app.post("/employees/edit/{employee_id}")
def update_employee(
    employee_id: int,
    name: str = Form(...),
    monthly_salary: float = Form(...),
    salary_cycle_start_day: int = Form(...),
    is_active: str | None = Form(None),
    db: Session = Depends(get_db),
    auth=Depends(login_required)
):
    if auth:
        return auth

    employee = db.get(Employee, employee_id)
    if not employee:
        return RedirectResponse("/employees", status_code=303)

    employee.name = name
    employee.monthly_salary = monthly_salary
    employee.per_day_salary = monthly_salary / 26
    employee.salary_cycle_start_day = min(salary_cycle_start_day, 28)
    employee.is_active = True if is_active else False
    db.commit()

    return RedirectResponse("/employees", status_code=303)

# ---------------- PENALTIES ----------------

@app.get("/penalties", response_class=HTMLResponse)
def penalties(
    request: Request,
    db: Session = Depends(get_db),
    auth=Depends(login_required)
):
    if auth:
        return auth

    return templates.TemplateResponse(
        "penalties.html",
        {
            "request": request,
            "employees": db.query(Employee).all(),
            "penalties": db.query(Penalty).order_by(Penalty.date.desc()).all()
        }
    )


@app.post("/penalties")
def add_penalty(
    employee_id: int = Form(...),
    reason: str = Form(...),
    amount: float = Form(...),
    db: Session = Depends(get_db),
    auth=Depends(login_required)
):
    if auth:
        return auth

    db.add(Penalty(
        employee_id=employee_id,
        date=date.today(),
        reason=reason,
        amount=amount
    ))
    db.commit()

    return RedirectResponse("/penalties", status_code=303)

# ---------------- ATTENDANCE UPLOAD ----------------

def iterate_rows(sheet):
    if hasattr(sheet, "iter_rows"):
        yield from sheet.iter_rows(values_only=True)
    else:
        for i in range(sheet.nrows):
            yield sheet.row_values(i)


@app.get("/attendance/upload", response_class=HTMLResponse)
def upload_attendance_form(
    request: Request,
    auth=Depends(login_required)
):
    if auth:
        return auth
    return templates.TemplateResponse("upload_attendance.html", {"request": request})


@app.post("/attendance/upload", response_class=HTMLResponse)
def upload_attendance(
    request: Request,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    auth=Depends(login_required)
):
    if auth:
        return auth

    filename = (file.filename or "").lower()

    if not filename.endswith((".xls", ".xlsx")):
        return templates.TemplateResponse(
            "upload_attendance.html",
            {"request": request, "error": "Unsupported file format"}
        )

    if filename.endswith(".xlsx"):
        wb = load_workbook(file.file, data_only=True)
        sheet = wb.active
    else:
        import xlrd
        book = xlrd.open_workbook(file_contents=file.file.read())
        sheet = book.sheet_by_index(0)

    rows = list(iterate_rows(sheet))
    emp_code = None
    header_row = None
    date_col = None
    in_cols, out_cols = [], []

    for i, row in enumerate(rows):
        for cell in row:
            if isinstance(cell, str) and "Emp.Code" in cell:
                m = re.search(r"(\d+)", cell)
                if m:
                    emp_code = m.group(1).zfill(8)

        if row and "Date" in row:
            header_row = i
            for idx, col in enumerate(row):
                if col == "Date":
                    date_col = idx
                elif isinstance(col, str) and col.startswith("In"):
                    in_cols.append(idx)
                elif isinstance(col, str) and col.startswith("Out"):
                    out_cols.append(idx)
            break

    if not emp_code or header_row is None or date_col is None:
        return templates.TemplateResponse(
            "upload_attendance.html",
            {"request": request, "error": "Invalid attendance format"}
        )

    employee = db.query(Employee).filter(
        Employee.emp_code == emp_code,
        (Employee.is_active == True) | (Employee.is_active.is_(None))
    ).first()

    if not employee:
        return templates.TemplateResponse(
            "upload_attendance.html",
            {"request": request, "error": "Employee not found"}
        )

    added = skipped = 0

    for row in rows[header_row + 1:]:
        if not row or not row[date_col]:
            continue

        work_date = row[date_col]
        if isinstance(work_date, datetime):
            work_date = work_date.date()

        if db.query(Attendance).filter_by(
            employee_id=employee.id,
            date=work_date
        ).first():
            skipped += 1
            continue

        in_time = next((row[i] for i in in_cols if i < len(row) and row[i]), None)
        out_time = next((row[i] for i in reversed(out_cols) if i < len(row) and row[i]), None)

        payable = 1 if in_time or out_time else 0
        status = "Present" if payable else "Absent"

        if in_time and isinstance(in_time, time) and in_time > time(10, 10):
            payable, status = 0.5, "Half"

        db.add(Attendance(
            employee_id=employee.id,
            date=work_date,
            in_time=in_time if isinstance(in_time, time) else None,
            out_time=out_time if isinstance(out_time, time) else None,
            status=status,
            payable_value=payable
        ))
        added += 1

    db.commit()

    return templates.TemplateResponse(
        "upload_attendance.html",
        {
            "request": request,
            "success": True,
            "employee_name": employee.name,
            "emp_code": employee.emp_code,
            "added": added,
            "skipped": skipped
        }
    )

# ---------------- SALARY ----------------

def calculate_salary(employee, db, start, end):
    records = db.query(Attendance).filter(
        Attendance.employee_id == employee.id,
        Attendance.date.between(start, end)
    ).all()

    present = sum(1 for r in records if r.payable_value == 1)
    half = sum(1 for r in records if r.payable_value == 0.5)
    absent = sum(1 for r in records if r.payable_value == 0)

    penalties = db.query(Penalty).filter(
        Penalty.employee_id == employee.id,
        Penalty.date.between(start, end)
    ).all()

    penalty_amount = sum(p.amount for p in penalties)
    payable_days = present + (half * 0.5)
    gross = payable_days * employee.per_day_salary
    net = gross - penalty_amount

    return present, half, absent, penalty_amount, penalties, gross, net


@app.get("/salary/generate", response_class=HTMLResponse)
def salary_generate_form(
    request: Request,
    db: Session = Depends(get_db),
    auth=Depends(login_required)
):
    if auth:
        return auth

    employees = db.query(Employee).filter(Employee.is_active == True).all()
    return templates.TemplateResponse(
        "salary_generate.html",
        {"request": request, "employees": employees}
    )


@app.post("/salary/generate", response_class=HTMLResponse)
def generate_salary(
    request: Request,
    employee_id: int = Form(...),
    month: int = Form(...),
    year: int = Form(...),
    db: Session = Depends(get_db),
    auth=Depends(login_required)
):
    if auth:
        return auth

    employee = db.get(Employee, employee_id)
    if not employee:
        return RedirectResponse("/salary/generate", status_code=303)

    start_day = employee.salary_cycle_start_day

    # ✅ January-safe cycle calculation
    if month == 1:
        cycle_start = date(year - 1, 12, start_day)
    else:
        cycle_start = date(year, month - 1, start_day)

    cycle_end = date(year, month, start_day) - timedelta(days=1)

    present, half, absent, penalty_amount, penalties, gross, net = (
        calculate_salary(employee, db, cycle_start, cycle_end)
    )

    run = SalaryRun(
        employee_id=employee.id,
        cycle_start=cycle_start,
        cycle_end=cycle_end,
        present_days=present,
        half_days=half,
        absent_days=absent,
        penalty_amount=penalty_amount,
        gross_salary=gross,
        net_salary=net
    )

    db.add(run)
    db.commit()

    return templates.TemplateResponse(
        "salary_result.html",
        {
            "request": request,
            "employee": employee,
            "cycle_start": cycle_start,
            "cycle_end": cycle_end,
            "present": present,
            "half": half,
            "absent": absent,
            "penalty_amount": penalty_amount,
            "penalties": penalties,
            "gross_salary": gross,
            "net_salary": net,
            "salary_run_id": run.id
        }
    )


@app.get("/salary/slip/{run_id}")
def download_salary_slip(
    run_id: int,
    db: Session = Depends(get_db),
    auth=Depends(login_required)
):
    if auth:
        return auth

    run = db.get(SalaryRun, run_id)
    if not run:
        return RedirectResponse("/", status_code=303)

    employee = db.get(Employee, run.employee_id)
    if not employee:
        return RedirectResponse("/", status_code=303)

    os.makedirs("temp", exist_ok=True)
    path = f"temp/salary_slip_{run.id}.pdf"

    c = canvas.Canvas(path, pagesize=A4)
    width, height = A4

    # ================= HEADER =================
    c.setFont("Helvetica-Bold", 16)
    c.drawCentredString(width / 2, height - 40, "SURJEET INTERNATIONAL")

    c.setFont("Helvetica", 9)
    c.drawCentredString(
        width / 2,
        height - 60,
        "D4 BASEMENT, SHUBHAM ENCLAVE, PASCHIM VIHAR, NEW DELHI - 110063"
    )
    c.drawCentredString(
        width / 2,
        height - 75,
        "Phone: 9911111013"
    )

    c.line(40, height - 90, width - 40, height - 90)

    # ================= TITLE =================
    c.setFont("Helvetica-Bold", 15)
    c.drawCentredString(width / 2, height - 115, "SALARY SLIP")

    # ================= EMPLOYEE INFO =================
    y = height - 150
    c.setFont("Helvetica-Bold", 10)
    c.drawString(50, y, "Employee Name:")
    c.drawString(300, y, "Employee Code:")

    c.setFont("Helvetica", 10)
    c.drawString(150, y, employee.name)
    c.drawString(420, y, employee.emp_code)

    y -= 20
    c.setFont("Helvetica-Bold", 10)
    c.drawString(50, y, "Pay Period:")
    c.setFont("Helvetica", 10)
    c.drawString(150, y, f"{run.cycle_start} to {run.cycle_end}")

    # ================= ATTENDANCE SUMMARY =================
    y -= 40
    c.setFont("Helvetica-Bold", 11)
    c.drawString(50, y, "ATTENDANCE SUMMARY")
    c.line(40, y - 5, width - 40, y - 5)

    y -= 25
    c.setFont("Helvetica", 10)
    c.drawString(50, y, f"Present Days : {run.present_days}")
    c.drawString(250, y, f"Half Days : {run.half_days}")

    payable_days = run.present_days + (run.half_days * 0.5)
    y -= 18
    c.drawString(50, y, f"Payable Days : {payable_days}")

    # ================= EARNINGS =================
    y -= 35
    c.setFont("Helvetica-Bold", 11)
    c.drawString(50, y, "EARNINGS")
    c.drawString(420, y, "AMOUNT (₹)")
    c.line(40, y - 5, width - 40, y - 5)

    y -= 25
    c.setFont("Helvetica", 10)
    c.drawString(50, y, "Salary for Payable Days")
    c.drawRightString(500, y, f"{run.gross_salary:,.2f}")

    # ================= DEDUCTIONS =================
    y -= 40
    c.setFont("Helvetica-Bold", 11)
    c.drawString(50, y, "DEDUCTIONS")
    c.drawString(420, y, "AMOUNT (₹)")
    c.line(40, y - 5, width - 40, y - 5)

    y -= 25
    c.setFont("Helvetica", 10)
    c.drawString(50, y, "Penalties")
    c.drawRightString(500, y, f"{run.penalty_amount:,.2f}")

    # ================= NET PAY =================
    y -= 50
    c.setFont("Helvetica-Bold", 13)
    c.rect(40, y - 15, width - 80, 35)
    c.drawCentredString(
        width / 2,
        y + 5,
        f"NET PAY : ₹ {run.net_salary:,.2f}"
    )

    # ================= FOOTER =================
    c.setFont("Helvetica", 8)
    c.setFillColorRGB(0.5, 0.5, 0.5)
    c.drawCentredString(
        width / 2,
        50,
        "This is a system-generated salary slip and does not require a signature."
    )

    c.showPage()
    c.save()

    return FileResponse(
        path,
        media_type="application/pdf",
        headers={
            "Content-Disposition": f'inline; filename=\"salary_slip_{run.id}.pdf\"'
        }
    )